#!/usr/bin/env python3
"""Build dist/forkable-sandboxes-cambridge-short.pptx for PowerPoint Presenter View.

One PowerPoint slide per page of talk-short.pdf (click steps included), each a full-bleed
render of that page. Its speaker notes hold only the script lines for that step, with
[CLICK], [ASK] and [PAUSE] cues and what comes next. Needs python-pptx and pdftoppm.
"""
import re, subprocess, tempfile
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt
from speaking_script import ROOT, deck, lines, mmss

TEAL, GOLD, MUTED, INK = (RGBColor(0x00, 0x7E, 0x87), RGBColor(0xAD, 0x70, 0x1B),
                          RGBColor(0x53, 0x68, 0x74), RGBColor(0x14, 0x2B, 0x3B))

def run(p, text, size=18, bold=False, color=INK):
    r = p.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color

def main():
    frames = []
    for f in deck():
        tex = (ROOT / 'short' / f"{f['name']}.tex").read_text()
        steps = max([int(n) for n in re.findall(r'<(\d+)', tex)] + [1])
        # Split the script at its click cues: segment k is spoken while step k+1 is on screen.
        segs = [[]]
        for marks, s, pause in lines(f):
            if 'CLICK' in marks: segs.append([])
            segs[-1].append((marks, s, pause))
        if len(segs) != steps: raise SystemExit(f"{f['name']}: {len(segs)} script segments for {steps} build steps")
        frames.append((f, segs))
    tmp = Path(tempfile.mkdtemp())
    subprocess.run(['pdftoppm', '-png', '-scale-to-x', '3200', '-scale-to-y', '-1',
                    str(ROOT / 'talk-short.pdf'), str(tmp / 'p')], check=True)
    pages = sorted(tmp.glob('p-*.png'))
    if len(pages) != sum(len(s) for _, s in frames): raise SystemExit('page count does not match build steps')
    prs = Presentation(); prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
    blank = prs.slide_layouts[6]; page = 0
    for i, (f, segs) in enumerate(frames):
        for k, seg in enumerate(segs):
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(pages[page]), 0, 0, prs.slide_width, prs.slide_height); page += 1
            tf = slide.notes_slide.notes_text_frame; tf.clear()
            head = tf.paragraphs[0]
            run(head, f"{i + 1} / {len(frames)}  ", 14, True, TEAL)
            run(head, f"{mmss(f['dur'])} planned · starts {mmss(f['start'])}", 14, False, MUTED)
            if len(segs) > 1: run(head, f"  ·  step {k + 1} of {len(segs)}", 14, True, TEAL)
            for marks, s, pause in seg:
                p = tf.add_paragraph(); p.space_before = Pt(6)
                for m in marks: run(p, f'[{m}] ', 18, True, TEAL if m == 'CLICK' else GOLD)
                run(p, s, 20)
                if pause: run(p, '  [PAUSE]', 16, False, MUTED)
            nxt = tf.add_paragraph(); nxt.space_before = Pt(12)
            if k + 1 < len(segs): run(nxt, '▸ Next: CLICK on this slide', 14, True, TEAL)
            elif i + 1 < len(frames): run(nxt, f"▸ Next: {i + 2}. {frames[i + 1][0]['title']}", 14, False, MUTED)
            else: run(nxt, '▸ Last slide: leave it up for questions.', 14, False, MUTED)
    out = ROOT / 'dist' / 'forkable-sandboxes-cambridge-short.pptx'
    prs.save(out)
    print(f'{len(frames)} slides, {page} PowerPoint slides -> {out}')

if __name__ == '__main__': main()
